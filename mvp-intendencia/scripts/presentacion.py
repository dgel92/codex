"""Generate editable slides and matching PDFs from the real demo evidence."""
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'entregables'
CAP = OUT / 'capturas'
INK, TEAL, DARK, BG, LIME, MUTED, LINE = '212121','E72D07','212121','F6F6F6','FFE5DE','777777','DDDDDD'
FONT_DIR = Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('Deck', str(FONT_DIR / 'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DeckBold', str(FONT_DIR / 'DejaVuSans-Bold.ttf')))
pdfmetrics.registerFontFamily('Deck', normal='Deck', bold='DeckBold', italic='Deck', boldItalic='DeckBold')


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(960), Pt(540)
        self.pdf = canvas.Canvas(str(OUT / 'Presentacion-MVP.pdf'), pagesize=(960,540))
        self.pdf.setTitle('EasyStock · Presentación funcional del MVP de Intendencia v1')
        self.index = 0

    def rect(self, x, y, w, h, color, radius=False):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = .08
        self.pdf.setFillColor(colors.HexColor('#'+color))
        self.pdf.roundRect(x, 540-y-h, w, h, 7 if radius else 0, fill=1, stroke=0)

    def text(self, x, y, w, text, size=14, color=INK, bold=False, line_height=1.35):
        # Explicit wrapping gives both formats the same line breaks.
        font = 'DeckBold' if bold else 'Deck'
        lines = []
        for paragraph in text.split('\n'):
            line = ''
            for word in paragraph.split(' '):
                proposed = (line+' '+word).strip()
                if pdfmetrics.stringWidth(proposed, font, size) > w and line:
                    lines.append(line); line = word
                else:
                    line = proposed
            lines.append(line)
        tf = self.slide.shapes.add_textbox(Pt(x),Pt(y-2),Pt(w),Pt(max(40,len(lines)*size*line_height+8))).text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line; p.font.name='DejaVu Sans'; p.font.size=Pt(size); p.font.bold=bold
            p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(0); p.line_spacing=Pt(size*line_height)
        self.pdf.setFillColor(colors.HexColor('#'+color)); self.pdf.setFont(font,size)
        for i, line in enumerate(lines):
            self.pdf.drawString(x,540-y-size-i*size*line_height,line)

    def image(self, name, x, y, w, h):
        path = CAP / (name+'.png')
        iw, ih = Image.open(path).size
        scale=min(w/iw,h/ih); dw, dh=iw*scale, ih*scale
        self.rect(x-3,y-3,dw+6,dh+6,LINE,True)
        self.slide.shapes.add_picture(str(path),Pt(x),Pt(y),Pt(dw),Pt(dh))
        self.pdf.drawImage(str(path),x,540-y-dh,width=dw,height=dh)

    def new(self, label, title=None, notes='', dark=False):
        if self.index:
            self.pdf.showPage()
        self.index+=1
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(0,0,960,540,DARK if dark else BG)
        self.text(36,24,760,'EASYSTOCK  /  '+label,10,LIME if dark else TEAL,True)
        if title:
            self.text(36,62,890,title,29,'FFFFFF' if dark else INK,True)
        self.rect(36,507,888,1,'444444' if dark else LINE)
        self.text(36,517,800,'FULLSTACKERS · INTENDENCIA · MVP ACADÉMICO · DATOS FICTICIOS',8,'AAAAAA' if dark else MUTED)
        self.text(900,517,28,f'{self.index:02}',8,'AAAAAA' if dark else MUTED)
        self.slide.notes_slide.notes_text_frame.text=notes

    def demo(self, hu, title, description, points, screenshot, notes):
        self.new(hu,title,notes)
        self.text(36,135,260,description,18,INK,True)
        y=230
        for point in points:
            self.rect(36,y+6,5,5,TEAL)
            self.text(51,y,238,point,12,MUTED)
            y+=64
        self.image(screenshot,326,113,598,385)

    def save(self):
        self.prs.save(str(OUT / 'Presentacion-MVP.pptx'))
        self.pdf.save()


def build_deck(report, total_tests):
    d=Deck()
    d.new('VERSIÓN 1.0',notes='Presentar el producto como implementación académica nueva. Las pantallas son capturas reales del navegador. No se usaron datos corporativos ni se modificó el ERP.',dark=True)
    d.text(36,101,370,'El pañol,\nen orden.',45,'FFFFFF',True,1.1)
    d.text(38,234,345,'EasyStock · Sistema de control de insumos de Intendencia.',18,'CCCCCC')
    d.rect(38,333,266,35,LIME,True);d.text(52,342,245,'MVP FUNCIONAL · HU-01 A HU-06',10,DARK,True)
    d.text(38,394,300,'Catálogo, ubicaciones y movimientos en una aplicación web.',12,'BBBBBB')
    d.image('02-inventario',410,103,514,372)

    d.demo('REFERENCIA VISUAL','La identidad de EasyStock, aplicada al pañol.',
           'MVP independiente con la estética de la página del equipo.',
           ['Logo y recursos visuales de FullStackers.', 'Rojo, gris oscuro y el acceso de EasyStock.', 'Archivos del equipo conservados sin cambios.'],
           '00-inicio', 'La entrega se agrega únicamente bajo mvp-intendencia/ en daniel-gonzalez. La estética se toma como referencia; no se cambia ni se integra el Angular o el Django existente del equipo.')

    d.new('EL PROBLEMA','Del retiro informal al registro compartido.',notes='El PDF describe consumos menores sin registro, insumos guardados sin identificación y diferencias entre stock físico y registrado. Esta entrega implementa el ciclo básico; todavía no mide mejoras en la empresa.')
    for x,num,title,body in [(36,'01','Registrar','Cada ingreso y cada consumo dejan un registro con usuario, cantidad y fecha.'),(338,'02','Encontrar','La búsqueda reúne disponibilidad, categoría y ubicación física en una sola vista.'),(640,'03','Confiar','Las salidas se validan contra el saldo y se confirman con una transacción atómica.')]:
        d.rect(x,153,284,259,'FFFFFF',True);d.text(x+23,177,220,num,33,TEAL)
        d.text(x+23,238,238,title,21,INK,True);d.text(x+23,287,235,body,14,MUTED)
    d.text(36,445,870,'Objetivo: apoyar el control interno. La mejora operativa requiere validación con usuarios reales.',12,MUTED)

    d.new('ALCANCE','Seis historias, un ciclo completo de inventario.',notes='La Etapa 10 del documento define HU-01 a HU-06 como Versión 1 y deja HU-07/HU-08 para una versión futura. Se persisten movimientos sin construir todavía el historial avanzado.')
    items=[('HU-01','Acceso por rol','Administrador y Operario'),('HU-02','Alta de insumos','Código único y ubicación'),('HU-03','Mínimos y máximos','Configuración autorizada'),('HU-04','Consulta de stock','Búsqueda y filtros combinados'),('HU-05','Consumo interno','Salida con sector y validación'),('HU-06','Recepción','Entrada con referencia opcional')]
    for i,(hu,title,body) in enumerate(items):
        x=36+(i%3)*302;y=135+(i//3)*155
        d.rect(x,y,284,135,'FFFFFF',True);d.text(x+18,y+17,246,hu,10,TEAL,True)
        d.text(x+18,y+46,248,title,17,INK,True);d.text(x+18,y+84,248,body,12,MUTED)
    d.text(36,465,880,'Versión 2: alertas visuales HU-07 e historial avanzado HU-08. Sin integración al ERP.',12,MUTED)

    d.new('DECISIONES TÉCNICAS','Simple de ejecutar. Con datos persistentes.',notes='Flask sirve las vistas y aplica la autorización. SQLite almacena usuarios, insumos, cambios de umbral y movimientos. Cantidades en milésimas enteras; BEGIN IMMEDIATE serializa validación y descuento. No se usa almacenamiento de navegador como base de datos.')
    for x,title,body in [(36,'Navegador','Interfaz responsive\nFormularios en español'),(348,'Python + Flask','Sesión y permisos\nValidaciones de negocio'),(660,'SQLite','Usuarios e insumos\nMovimientos y umbrales')]:
        d.rect(x,150,264,131,'FFFFFF',True);d.text(x+20,171,225,title,20,INK,True);d.text(x+20,216,224,body,12,MUTED)
    d.text(308,201,32,'→',22,TEAL);d.text(620,201,32,'→',22,TEAL)
    d.rect(36,318,888,134,'FFF0EB',True)
    d.text(58,338,835,'Reglas verificadas',17,INK,True)
    d.text(58,376,820,'Contraseñas con hash scrypt · Permisos en servidor · Protección CSRF\nStock sin redondeo flotante · Descuento atómico · Reenvíos sin duplicación',14,INK)

    d.demo('HU-01','Cada rol tiene su espacio.','La sesión identifica al responsable de cada operación.',['Administrador: alta y umbrales.','Operario: búsqueda y movimientos.','Las rutas también verifican permisos.'],'01-login','Mostrar las dos cuentas ficticias documentadas. No presentar las cuentas demo como usuarios corporativos. La autorización del servidor devuelve 403 al Operario que solicita el alta.')
    d.demo('HU-02','Primero, identificar el material.','Desengrasante industrial, guardado en Puerta 3.',['Código y nombre únicos.','Categoría y unidad obligatorias.','El alta comienza con stock cero.'],'03-alta','En el video se crea LIM-003, Desengrasante industrial, Limpieza, litro, Puerta 3. La entrada formal cargará las existencias.')
    d.demo('HU-03','Definir los niveles de referencia.','Umbrales editables por el Administrador.',['Valores positivos; mínimo menor al máximo.','Fecha y usuario guardados internamente.','Las alertas pertenecen a la Versión 2.'],'04-umbrales','La demostración cambia mínimo 3/máximo 25 a 4/30. Para unidad se admiten enteros; para litros y kg, tres decimales. No se bloquea una recepción por superar el máximo.')
    d.demo('HU-06','Registrar lo que llega al depósito.','Una recepción de 10 litros inicia el ciclo.',['Referencia: Remito DEMO-001.','Usuario y fecha automáticos.','El stock pasa de 0 a 10 litros.'],'05-entrada','Como Operario registrar una entrada de 10 litros. La referencia es opcional. Se guarda cantidad y responsable en SQLite.')
    d.demo('HU-04','Encontrar sin recorrer el pañol.','Nombre, categoría y ubicación se pueden combinar.',['Búsqueda parcial de texto.','Cantidad y unidad visibles.','Actualizar consulta el saldo vigente.'],'06-busqueda','Buscar desengrasante, filtrar Limpieza y Puerta 3. Explicar que la aplicación refresca en cada consulta y al confirmar movimientos propios; no implementa actualización push entre dispositivos.')
    d.demo('HU-05','Un retiro nunca puede dejar stock negativo.','12 litros solicitados. Solo 10 disponibles.',['El sistema informa stock insuficiente.','La salida no crea un movimiento.','El saldo permanece intacto.'],'07-stock-insuficiente','Mostrar el error real. Las pruebas también ejecutan dos retiros simultáneos de 20 litros frente a 24 disponibles; solo uno puede completarse.')
    d.demo('HU-05 / RESULTADO','Un consumo registrado, un saldo comprobable.','10 − 2,5 = 7,5 litros disponibles.',['Destino: sector Molino.','Responsable: Operario demo.','Saldo confirmado en pantalla y SQL.'],'09-resultado','Corregir la cantidad a 2,5 litros y confirmar. La verificación independiente en SQLite comprueba una entrada de 10000 milésimas y una salida de 2500 milésimas, ambas del Operario.')

    d.new('VALIDACIÓN','Evidencias ejecutadas en esta entrega.',notes='Estos resultados pertenecen a esta sesión de implementación, no a la instancia académica anterior. Las pruebas automatizadas no equivalen a aceptación de la organización. Los archivos de salida acompañan la entrega.')
    for x,num,title in [(36,str(total_tests),'pruebas unitarias e integrales'),(338,str(report['aprobados']),'verificaciones del recorrido real'),(640,'0','fallos en las ejecuciones finales')]:
        d.rect(x,137,284,132,'FFFFFF',True);d.text(x+22,149,240,num,46,TEAL,True);d.text(x+22,221,245,title,12,MUTED)
    rows=[('Reglas de negocio','Duplicados, cantidades y umbrales'),('Integridad del stock','Stock insuficiente y retiros concurrentes'),('Acceso y persistencia','Roles, CSRF, reinicio y reenvío de formularios'),('Demo funcional','Navegación real en Chromium y comprobación SQL')]
    for i,(a,b) in enumerate(rows):
        y=302+i*39;d.text(39,y,255,a,12,INK,True);d.text(312,y,608,b,12,MUTED)
    d.text(36,471,878,'Registros adjuntos: pruebas-unitarias-integracion.txt y pruebas-navegador.json.',10,TEAL)

    d.new('ENTREGA Y PRÓXIMOS PASOS','Versión 1 construida. Validación operativa pendiente.',notes='La entrega incluye código, instrucciones, presentación editable y PDF, video y evidencias. El MVP funciona localmente con datos ficticios. Faltan validación por usuarios de la empresa y preparación productiva. No afirmar publicación o despliegue.')
    d.rect(36,137,429,310,'FFFFFF',True);d.rect(485,137,439,310,'FFF0EB',True)
    d.text(58,157,383,'Disponible para la exposición',19,INK,True)
    d.text(58,208,378,'MVP separado en mvp-intendencia/\nDiseño de EasyStock del equipo\nManual de usuario y guion de demo\nVideo MP4 subtitulado, sin audio\nPresentación PowerPoint y PDF\nInforme y registros de pruebas',14,MUTED,False,2)
    d.text(508,157,391,'Siguiente instancia',19,INK,True)
    d.text(508,208,387,'Validar roles y ubicaciones con el sector.\nProbar el flujo con usuarios reales.\nPreparar usuarios, respaldos y despliegue.\nIterar sobre HU-07 y HU-08.',14,INK,False,2)
    d.text(36,469,888,'Demo adjunta: Demo-MVP.mp4 · Instrucciones reproducibles: README.md',11,TEAL)
    d.save()
    return d.index


def build_report(report, total_tests):
    styles=getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName='Deck';style.textColor=colors.HexColor('#'+INK)
    styles['Title'].fontName='DeckBold';styles['Heading1'].fontName='DeckBold';styles['Heading2'].fontName='DeckBold'
    styles['BodyText'].fontSize=9;styles['BodyText'].leading=13
    doc=SimpleDocTemplate(str(OUT / 'Informe-de-validacion.pdf'),pagesize=(595,842),rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=40)
    story=[]
    def p(text, style='BodyText'):
        story.append(Paragraph(text,styles[style]));story.append(Spacer(1,6))
    p('EasyStock · MVP de Intendencia v1', 'Title')
    p('Informe de implementación y validación', 'Heading1')
    local=datetime.fromisoformat(report['ejecutado_en_utc']).astimezone(timezone(timedelta(hours=-3)))
    p('Ejecución registrada: '+local.strftime('%d/%m/%Y %H:%M')+' (Argentina, UTC−3).')
    p('Fuente de requisitos: ABP- PP (1).pdf, en especial las historias de usuario de la etapa 9 y el roadmap de la etapa 10. Alcance de esta implementación: HU-01 a HU-06. Datos y cuentas ficticios.')
    p('Resultado', 'Heading2')
    p(f'{total_tests} pruebas unitarias e integrales aprobadas, sin fallos ni omisiones en esa suite. {report["aprobados"]} verificaciones funcionales aprobadas en Chromium, incluida comprobación de movimientos en SQLite. El video adjunto registra este recorrido real; no tiene audio y contiene subtítulos en español.')
    p('El diseño toma como referencia la página de EasyStock del equipo FullStackers. Se agrega exclusivamente una carpeta independiente mvp-intendencia/ en daniel-gonzalez. Los archivos existentes del equipo permanecen sin cambios.')
    p('Cobertura funcional', 'Heading2')
    data=[['Historia','Capacidad implementada'],['HU-01','Autenticación, sesión y autorización por rol'],['HU-02','Alta con unicidad, ubicación estructurada y stock cero'],['HU-03','Umbrales válidos y registro de administrador/fecha'],['HU-04','Búsqueda parcial y filtros combinados'],['HU-05','Salidas atómicas, sector y bloqueo por falta de stock'],['HU-06','Entradas y trazabilidad de recepción']]
    table=Table(data,colWidths=[70,441]);table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'Deck'),('FONTSIZE',(0,0),(-1,-1),9),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+LIME)),('BOTTOMPADDING',(0,0),(-1,-1),10),('TOPPADDING',(0,0),(-1,-1),10),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#'+LINE))]));story.append(table)
    p('Integridad y acceso', 'Heading2')
    p('La suite verifica cantidades positivas, precisión decimal y unidades enteras; unicidad de código/nombre; permisos tanto en interfaz como en servidor; protección CSRF; salidas sin saldo; reenvíos idempotentes; dos retiros concurrentes; persistencia y carga demo repetible. Se utilizan bases temporales aisladas.')
    story.append(PageBreak())
    p('Recorrido funcional ejecutado', 'Heading1')
    for i,case in enumerate(report['casos'],1):
        p(f'{i:02}. {case["caso"]} — {case["resultado"]}')
    p('Operaciones de la demostración', 'Heading2')
    p('Se crea LIM-003 (Desengrasante industrial, litro, Puerta 3) con stock cero. Como Operario se ingresan 10 litros, se rechaza un retiro de 12 y se confirma una salida de 2,5 al sector Molino. El saldo final es 7,5 litros. La consulta SQL independiente comprueba exactamente dos movimientos del nuevo insumo y el usuario responsable.')
    p('Archivos de evidencia', 'Heading2')
    p('pruebas-unitarias-integracion.txt: salida completa del runner unittest.<br/>pruebas-navegador.json: verificaciones y datos de los movimientos.<br/>Demo-MVP.mp4 y Demo-MVP.srt: grabación y subtítulos.<br/>capturas/: pantallas reales usadas en las diapositivas.<br/>tests/test_mvp.py y scripts/demo.py: procedimientos reproducibles.')
    p('Límites de la validación', 'Heading2')
    p('Son pruebas técnicas automatizadas realizadas durante esta sesión de desarrollo. No constituyen validación con usuarios de la empresa, pruebas de carga, auditoría de seguridad, medición de mejoras ni despliegue en producción. Tampoco prueban ejecución en la instancia académica anterior. HU-07 y HU-08 se mantienen fuera de esta versión. La consulta de stock se actualiza por petición; no hay notificaciones push.')
    p('Reproducción', 'Heading2')
    p('Seguir README.md para instalar requirements.txt e iniciar app.py --seed-demo. Ejecutar python -m unittest discover -s tests -v. Con requirements-dev.txt, Chromium y ffmpeg, ejecutar scripts/demo.py y scripts/presentacion.py. El grabador usa una base temporal y detiene su propio servidor.')
    doc.build(story)


def manual_pdf():
    styles=getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName='Deck';style.fontSize=10;style.leading=15
    styles['Title'].fontName='DeckBold';styles['Title'].fontSize=23;styles['Title'].leading=30
    styles['Heading2'].fontName='DeckBold';styles['Heading2'].fontSize=14;styles['Heading2'].leading=20
    story=[]
    for paragraph in (ROOT/'docs/MANUAL.md').read_text().split('\n\n'):
        style='BodyText'
        if paragraph.startswith('# '):style='Title';paragraph=paragraph[2:]
        elif paragraph.startswith('## '):style='Heading2';paragraph=paragraph[3:]
        paragraph=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',paragraph)
        story.extend([Paragraph(paragraph.replace('\n','<br/>'),styles[style]),Spacer(1,11)])
    SimpleDocTemplate(str(OUT/'Manual-de-usuario.pdf'),rightMargin=45,leftMargin=45,topMargin=40,bottomMargin=40).build(story)


if __name__ == '__main__':
    report=json.loads((OUT/'pruebas-navegador.json').read_text())
    log=(OUT/'pruebas-unitarias-integracion.txt').read_text()
    total=int(re.search(r'Ran (\d+) tests',log).group(1))
    assert '\nOK\n' in log and report['fallidos']==0
    count=build_deck(report,total)
    build_report(report,total)
    manual_pdf()
    print(f'{count} diapositivas; PowerPoint, PDF, informe de validación y manual generados.')
